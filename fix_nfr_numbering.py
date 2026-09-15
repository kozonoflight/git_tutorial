#!/usr/bin/env python3
"""(1/2)ページで番号を共有する方式に合わせて通し番号と相互参照を直す。

番号は 1.SLA / 2.サービス対応 / 3.全体俯瞰 / 4.対応内容一覧 / 5.監視 /
6.バックアップ / 7.セキュリティ / 8.システム構成 の8つ。
枝番ページは同じ番号を共有するため、参照は 8.（1/3）の形式にする。
"""

import copy
import sys

from pptx import Presentation

SRC = sys.argv[1] if len(sys.argv) > 1 else "/home/ubuntu/.cursor/projects/workspace/uploads/____________20260915_-_____3754.pptx"
DST = sys.argv[2] if len(sys.argv) > 2 else "/workspace/要件定義書_非機能要件_20260917_v3.pptx"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def iter_text_frames(slide):
    for shape in slide.shapes:
        if shape.has_text_frame:
            yield shape.text_frame
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame
        if shape.shape_type == 6:
            yield from _group_frames(shape)


def _group_frames(group):
    for sub in group.shapes:
        if sub.shape_type == 6:
            yield from _group_frames(sub)
        elif sub.has_text_frame:
            yield sub.text_frame


def _tf(obj):
    return obj.text_frame if hasattr(obj, "text_frame") else obj


def replace_in_frame(tf, old, new):
    hit = 0
    for para in tf.paragraphs:
        runs = para.runs
        if not runs:
            continue
        joined = "".join(r.text for r in runs)
        if old not in joined:
            continue
        runs[0].text = joined.replace(old, new)
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        hit += 1
    return hit


def replace_on_slide(slide, old, new, label="", expect=1):
    hit = sum(replace_in_frame(tf, old, new) for tf in iter_text_frames(slide))
    if hit != expect:
        print(f"  !! {label}: 期待{expect}件 / 実際{hit}件  {old[:34]!r}")
    return hit


def set_lines(obj, lines):
    tf = _tf(obj)
    base = next((p for p in tf.paragraphs if p.runs), None)
    if base is None:
        return False
    base_xml = copy.deepcopy(base._p)
    body = tf._txBody
    for p in list(tf.paragraphs):
        body.remove(p._p)
    for line in lines:
        el = copy.deepcopy(base_xml)
        runs = el.findall(f"{A_NS}r")
        for extra in runs[1:]:
            el.remove(extra)
        for br in el.findall(f"{A_NS}br"):
            el.remove(br)
        runs[0].find(f"{A_NS}t").text = line
        body.append(el)
    return True


def append_line(obj, line):
    tf = _tf(obj)
    base = None
    for p in tf.paragraphs:
        if p.runs:
            base = p
    if base is None:
        return False
    el = copy.deepcopy(base._p)
    runs = el.findall(f"{A_NS}r")
    for extra in runs[1:]:
        el.remove(extra)
    for br in el.findall(f"{A_NS}br"):
        el.remove(br)
    runs[0].find(f"{A_NS}t").text = line
    tf._txBody.append(el)
    return True


def find_title_shape(slide):
    for shape in slide.shapes:
        try:
            if shape.is_placeholder and shape.placeholder_format.type == 1:
                return shape
        except Exception:
            pass
    best = None
    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        if shape.top is None or shape.top > 500000:
            continue
        if best is None or shape.width > best.width:
            best = shape
    return best


prs = Presentation(SRC)
slides = list(prs.slides)
if len(slides) != 15:
    sys.exit(f"想定外のページ数: {len(slides)}")

# ============================================================
# 1. 通し番号（枝番ページは同番号を共有）。ダッシュは半角ハイフンに統一
# ============================================================
TITLES = {
    2: "1.SLA（1/2）",
    3: "1.SLA（2/2）",
    4: "2.サービス対応",
    5: "3.全体俯瞰",
    6: "4.対応内容一覧（1/2）",
    7: "4.対応内容一覧（2/2）",
    8: "5.対応内容説明 - 監視（システム構成）",
    9: "6.対応内容説明 - バックアップ",
    10: "7.対応内容説明 - セキュリティ（1/2）",
    11: "7.対応内容説明 - セキュリティ（2/2）",
    12: "8.対応内容説明 - システム構成（1/3）",
    13: "8.対応内容説明 - システム構成（2/3）",
    14: "8.対応内容説明 - システム構成（3/3）",
}
print("== タイトル ==")
for no, text in TITLES.items():
    shape = find_title_shape(slides[no - 1])
    if shape is None:
        print(f"  !! p{no} タイトル未検出")
        continue
    before = shape.text_frame.text.replace("\n", " ")
    if before == text:
        print(f"  p{no}: {text}  (変更なし)")
        continue
    set_lines(shape, [text])
    print(f"  p{no}: {before}  ->  {text}")

# ============================================================
# 2. 相互参照。枝番ページと区別するため 8.（1/3）形式にする
# ============================================================
print("== 相互参照 ==")
replace_on_slide(slides[2], "本案件の想定は11.に記載", "本案件の想定は8.（1/3）に記載", label="p3 同時接続")
replace_on_slide(slides[2], "（詳細は8.）", "（詳細は6.）", label="p3 保存期間")

s = slides[4]
replace_on_slide(s, "7.に記載", "5.に記載", label="p5 システム運用")
replace_on_slide(s, "8.に記載", "2.6.に記載", label="p5 保守")
replace_on_slide(s, "9.10.に記載", "7.に記載", label="p5 安全性・セキュリティ", expect=2)
replace_on_slide(s, "11.12.13.に記載", "8.に記載", label="p5 拡張性")
replace_on_slide(s, "11.12.に記載", "8.に記載", label="p5 性能", expect=0)

replace_on_slide(slides[11], "（12.の成長倍率2倍に対応）", "（8.（2/3）の成長倍率2倍に対応）", label="p12 成長軸")

s = slides[12]
replace_on_slide(s, "11.および本ページの結果より、", "前ページおよび本ページの結果より、", label="p13 結論")
replace_on_slide(s, "（11.のアクティブユーザ2倍に対応）", "（8.（1/3）のアクティブユーザ2倍に対応）", label="p13 成長軸")
replace_on_slide(s, "採用値は11.の116名", "採用値は8.（1/3）の116名", label="p13 採用値")

# ============================================================
# 3. 前版で直した内容の取り込み漏れ
# ============================================================
print("== 取り込み漏れ ==")
replace_on_slide(slides[5], "補完場所はクラウドサービス内", "保管場所はクラウドサービス内", label="p6 保管場所")
replace_on_slide(slides[9], "XSS,SSQLインジェクション", "XSS,SQLインジェクション", label="p10 SQL")

tbl = next(sh.table for sh in slides[1].shapes if sh.has_table)
for row in tbl.rows:
    if row.cells[1].text.strip() == "重大障害時の代替手段" and "24時間" not in row.cells[3].text:
        append_line(row.cells[3], "日次バックアップのため最大24時間分の戻りが発生")
        print("  p2 重大障害時の備考に戻り時間を追記")

# ============================================================
# 4. 項目列の重複ラベル
# ============================================================
tbl = next(sh.table for sh in slides[2].shapes if sh.has_table)
seen = set()
for ri, row in enumerate(tbl.rows):
    head = row.cells[0].text.strip()
    if not head:
        continue
    if head in seen:
        set_lines(row.cells[0], [""])
        print(f"  p3 r{ri} の重複ラベル「{head}」を空欄化")
    seen.add(head)

prs.save(DST)
print("\nwrote", DST)
