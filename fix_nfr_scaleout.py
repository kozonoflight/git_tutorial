#!/usr/bin/env python3
"""ハードウェア拡張の記載を実構成（LB1台＋WEB・DB1台）に合わせる。

標準構成はロードバランサ1台にWEB・DBサーバ1台。自動的なスケールアウトは
行わず、複数台化はオプション（有償）のインフラ拡張で対応する。
"""

import copy
import sys

from pptx import Presentation

SRC = sys.argv[1] if len(sys.argv) > 1 else "/workspace/要件定義書_非機能要件_20260917_v3.pptx"
DST = sys.argv[2] if len(sys.argv) > 2 else "/workspace/要件定義書_非機能要件_20260917_v4.pptx"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _tf(obj):
    return obj.text_frame if hasattr(obj, "text_frame") else obj


def set_lines(obj, lines, like=None):
    """内容を行単位で入れ替える。

    空セルにはランが無く複製元が作れないため、like に書式の見本となる
    セル/テキストフレームを渡せるようにする。
    """
    tf = _tf(obj)
    base = next((p for p in tf.paragraphs if p.runs), None)
    if base is None and like is not None:
        base = next((p for p in _tf(like).paragraphs if p.runs), None)
    if base is None:
        return False
    base_xml = copy.deepcopy(base._p)
    body = tf._txBody
    for p in list(tf.paragraphs):
        body.remove(p._p)
    for line in lines:
        el = copy.deepcopy(base_xml)
        runs = el.findall(f"{A_NS}r")
        for extra in runs[1:]:
            el.remove(extra)
        for br in el.findall(f"{A_NS}br"):
            el.remove(br)
        runs[0].find(f"{A_NS}t").text = line
        body.append(el)
    return True


def replace_in_frame(tf, old, new):
    hit = 0
    for para in tf.paragraphs:
        runs = para.runs
        if not runs:
            continue
        joined = "".join(r.text for r in runs)
        if old not in joined:
            continue
        runs[0].text = joined.replace(old, new)
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        hit += 1
    return hit


prs = Presentation(SRC)

# ============================================================
# 4.対応内容一覧（2/2） ハードウェア拡張
# ============================================================
print("== 4.対応内容一覧（2/2） ==")
tbl = next(sh.table for sh in prs.slides[6].shapes if sh.has_table)

SCALE_OUT = [
    "[内容]（標準対応）",
    "・LBサーバ1台、WEB・DBサーバ1台の構成とし、自動的なスケールアウトは行わない。",
    "・WEBサーバの複数台化、DBサーバの分離はオプション（有償）のインフラ拡張で対応する。",
]
MANUAL_NOTE = "自動スケールではなく、停止を伴う手動の構成変更となります。"

for ri, row in enumerate(tbl.rows):
    req = row.cells[1].text.strip()
    if req == "スケールアウト構成":
        before = row.cells[2].text.replace("\n", " / ")
        set_lines(row.cells[2], SCALE_OUT)
        print(f"  r{ri} スケールアウト構成")
        print(f"     旧: {before}")
        print(f"     新: {' / '.join(SCALE_OUT)}")
        set_lines(
            row.cells[3],
            ["初期構成は8.（3/3）に記載のとおりです。", "構成変更はオプション（有償）対応です。"],
        )
        print("     備考も更新")
    elif req in ("スケールアップの可能性", "動的なリソース拡張", "動的なリソース再編成"):
        hit = replace_in_frame(row.cells[2].text_frame, "・VPSの構成変更によって対応する。", "・クラウドサーバのプラン変更によって対応する。")
        filled = False
        if not row.cells[3].text.strip():
            filled = set_lines(row.cells[3], [MANUAL_NOTE], like=row.cells[2])
        print(f"  r{ri} {req}: 本文{hit}件置換 / 備考補完={filled}")

# ============================================================
# 8.対応内容説明 - システム構成（2/3） プラン2の位置づけ
# ============================================================
print("== 8.（2/3） ==")
for shape in prs.slides[12].shapes:
    if shape.has_text_frame and "営業日30日/月" in shape.text_frame.text:
        lines = [
            "".join(r.text for r in p.runs) for p in shape.text_frame.paragraphs if "".join(r.text for r in p.runs)
        ]
        lines.append("※ プラン2（拡張）はオプション（有償）のインフラ拡張となります（標準はプラン1の1台構成）。")
        set_lines(shape, lines)
        print("  プラン2がオプションである旨を注記に追加")
        break

prs.save(DST)
print("\nwrote", DST)
