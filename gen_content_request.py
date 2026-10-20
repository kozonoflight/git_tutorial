#!/usr/bin/env python3
"""サイト掲載コンテンツのご提供依頼Excelを作る。

フッターの定型ページは、現ひな型がシステム提供会社の情報・消費者向け文面のままなので、
掲載主体であるローリングスジャパン様に原稿を作っていただく。
"""

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

OUT = "/workspace/ご提供依頼_サイト掲載コンテンツ.xlsx"

NAVY = "1B3A4B"
TEAL = "1D7A70"
BLUE = "3D5A80"
WHITE = "FFFFFF"
INK = "1B3A4B"
GRAY = "5C6770"
PALE = "F4F7FB"
TEAL_LT = "E6F4F2"
YELLOW = "FFF4D6"
CREAM = "FFF8EE"
LINE = "D5DBE0"
RED = "C0392B"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
wrap = Alignment(wrap_text=True, vertical="center")
left = Alignment(wrap_text=True, vertical="center", horizontal="left")

wb = Workbook()


def font(size=11, bold=False, color=INK, name="メイリオ"):
    return Font(name=name, size=size, bold=bold, color=color)


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def paint(cell, value, size=10, bold=False, color=INK, bg=None, align=None):
    cell.value = value
    cell.font = font(size, bold, color)
    cell.alignment = align or left
    cell.border = thin
    if bg:
        cell.fill = fill(bg)
    return cell


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def page_setup(ws, title):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.5, header=0.25, footer=0.25)
    ws.oddHeader.left.text = title
    ws.oddFooter.right.text = "ご記入後、本ファイルをご返送ください"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.freeze_panes = "A8"
    ws.auto_filter.ref = None
    ws.sheet_properties.tabColor = TEAL
    ws.print_title_rows = "1:7"
    ws.page_setup.horizontalCentered = True
    ws.sheet_format.defaultRowHeight = 18


def banner(ws, last_col, page_name, what, why):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    c = ws.cell(1, 1, "サイト掲載コンテンツのご提供依頼")
    c.font = font(16, True, WHITE)
    c.fill = fill(NAVY)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    c = ws.cell(2, 1, page_name)
    c.font = font(14, True, TEAL)
    c.fill = fill(TEAL_LT)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 24

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    c = ws.cell(3, 1, "作っていただくもの：" + what)
    c.font = font(10, False, INK)
    c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    ws.row_dimensions[3].height = 32

    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=last_col)
    c = ws.cell(4, 1, "なぜ必要か：" + why)
    c.font = font(10, False, GRAY)
    c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    ws.row_dimensions[4].height = 32

    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=last_col)
    c = ws.cell(
        5,
        1,
        "黄色のセルをご記入ください。掲載しない項目は「掲載要否」を「掲載しない」にしてください。現ひな型はシステム提供会社の情報が入った見本であり、そのまま掲載しません。",
    )
    c.font = font(9, False, "7A4E00")
    c.fill = fill(CREAM)
    c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 8


def input_table(ws, start, headers, rows, widths, yesno_col=None):
    """headers: list. rows: list of tuples matching headers. None means input cell."""
    for col, h in enumerate(headers, 1):
        paint(ws.cell(start, col), h, 10, True, WHITE, NAVY, Alignment(wrap_text=True, vertical="center", horizontal="center"))
    ws.row_dimensions[start].height = 28
    for r, row in enumerate(rows):
        excel_r = start + 1 + r
        ws.row_dimensions[excel_r].height = 36
        for c, value in enumerate(row, 1):
            bg = YELLOW if value is None else (PALE if r % 2 else WHITE)
            paint(ws.cell(excel_r, c), "" if value is None else value, 10, False, INK, bg)
    set_widths(ws, widths)
    if yesno_col:
        dv = DataValidation(type="list", formula1='"掲載する,掲載しない,確認中"', allow_blank=True)
        dv.error = "掲載する / 掲載しない / 確認中 から選んでください"
        dv.errorTitle = "掲載要否"
        dv.add(f"{get_column_letter(yesno_col)}{start+1}:{get_column_letter(yesno_col)}{start+len(rows)}")
        ws.add_data_validation(dv)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start+len(rows)}"
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start+len(rows)}"
    return start + len(rows)


# ---------------------------------------------------------------------------
# 00 一覧
# ---------------------------------------------------------------------------
index = wb.active
index.title = "00_ご依頼一覧"

items = [
    ("1", "フッター共通", "01_フッター", "新規作成", "電話番号、受付時間、フッターに出すリンクの採否", "画面下部に常時表示するため、電話番号と受付時間は貴社の窓口情報である必要があります。", "高"),
    ("2", "会社概要", "02_会社概要", "新規作成", "サイトに掲載する貴社の会社情報", "現ひな型はシステム提供会社の会社概要です。注文サイトの運営主体として貴社情報へ差し替えます。", "高"),
    ("3", "お問い合わせ", "03_お問い合わせ", "新規作成", "問い合わせ窓口、受付方法、フォームに載せる項目", "注文・納期・返品の問い合わせ先が貴社CSになるため、宛先と振り分けを決めていただく必要があります。", "高"),
    ("4", "ご利用ガイド", "04_ご利用ガイド", "確認・加筆", "標準の操作説明に、貴社の注文ルールを足したガイド", "操作手順は標準機能で説明できます。通常／カスタム／展示会、与信、承認、キャンセルは貴社ルールの文章が必要です。", "高"),
    ("5", "ご利用規約", "05_ご利用規約", "確認・加筆", "サイト利用条件。特に注文の成立時点", "標準文面は消費者向け通販の前提が残っています。BtoBの成立時点（与信、承認）を貴社で確定していただく必要があります。", "高"),
    ("6", "よくあるご質問", "06_よくあるご質問", "新規作成", "取引先向けの質問と回答", "操作の一般論ではなく、価格の見え方、展示会、代理注文、返品など貴社固有の質問が中心になります。", "中"),
    ("7", "プライバシーポリシー", "07_プライバシーポリシー", "新規作成", "個人情報の取扱方針。空欄の社名、管理者、共同利用先", "現ひな型は社名・管理者・共同利用先が空欄です。取得する情報の範囲も貴社業務に合わせて確定が必要です。", "高"),
    ("8", "セキュリティポリシー", "08_セキュリティポリシー", "新規作成", "情報セキュリティ方針の本文", "現ひな型は見出しだけで本文が空欄です。方針を掲出するかどうかを含め、文章をご用意ください。", "中"),
    ("9", "特定商取引法に基づく表記", "09_特定商取引法", "要否判断＋原稿", "掲載する場合の販売業者、支払、引渡、返品条件", "現ひな型は消費者向け（クレジット、コンビニ、地域別送料）です。BtoBの取引条件とは別物なので、掲載要否からご判断ください。", "高"),
]

last = 6
banner(
    index,
    last,
    "掲載ページの原稿依頼",
    "フッターから遷移する各ページについて、掲載主体である貴社の原稿をこのブックに記入していただきます。",
    "システム側で文面を確定できないためです。会社情報、取引条件、個人情報の取扱い、問い合わせ窓口は貴社の決定事項です。",
)
headers = ["No", "掲載ページ", "記入シート", "依頼区分", "このシートで作っていただくもの", "優先"]
for col, h in enumerate(headers, 1):
    paint(index.cell(7, col), h, 10, True, WHITE, NAVY, Alignment(wrap_text=True, vertical="center", horizontal="center"))
index.row_dimensions[7].height = 24
for i, item in enumerate(items):
    r = 8 + i
    index.row_dimensions[r].height = 48
    vals = (item[0], item[1], item[2], item[3], item[4], item[6])
    for c, v in enumerate(vals, 1):
        bg = TEAL_LT if item[6] == "高" else WHITE
        paint(index.cell(r, c), v, 10, c == 1, INK, bg, Alignment(wrap_text=True, vertical="center", horizontal="center" if c in (1, 4, 6) else "left"))
    # なぜ必要かをコメントではなく別列にすると横長になる。5列に要旨、備考はセルコメント
    index.cell(r, 5).comment = Comment(item[5], "依頼趣旨")
    index.cell(r, 5).comment.width = 280
    index.cell(r, 5).comment.height = 80

index.merge_cells("A17:F17")
note = index.cell(17, 1, "記入のお願い：黄色セルが貴社記入欄です。各シート上部の「作っていただくもの」を読んだうえで記入してください。文章が長いページ（規約・ガイド・ポリシー）は、このExcelに全文を書いていただいても、別紙Wordを添付いただいても構いません。別紙の場合は該当セルに「別紙」と記載してください。")
note.font = font(9, False, GRAY)
note.alignment = Alignment(wrap_text=True, vertical="center")
index.row_dimensions[17].height = 42
set_widths(index, [8, 28, 24, 16, 62, 10])
index.freeze_panes = "A8"
index.auto_filter.ref = "A7:F16"
index.sheet_view.showGridLines = False
index.sheet_properties.tabColor = NAVY
index.page_setup.orientation = "landscape"
index.page_setup.fitToPage = True
index.page_setup.fitToWidth = 1
index.page_setup.fitToHeight = 1
index.page_setup.paperSize = index.PAPERSIZE_A4
index.sheet_properties.pageSetUpPr.fitToPage = True
index.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.4)
index.print_title_rows = "1:7"
index.oddHeader.left.text = "サイト掲載コンテンツのご提供依頼"
index.page_setup.horizontalCentered = True
index.sheet_view.zoomScale = 120

# ---------------------------------------------------------------------------
# 01 フッター
# ---------------------------------------------------------------------------
ws = wb.create_sheet("01_フッター")
banner(
    ws, 4,
    "フッター（全ページ共通）",
    "画面最下部に常時出す電話番号、受付時間、リンクの採否です。",
    "全画面の共通部品のため、ここに貴社以外の電話番号が残ると全ページに表示されます。",
)
input_table(
    ws, 7,
    ["項目", "掲載要否", "ご記入内容", "確認したいこと"],
    [
        ("ロゴ横のサイト名", "掲載する", None, "現ひな型は「B2B ORDER SITE」。貴社の呼び方（例：Rawlings B2B ORDER SITE）があれば記入"),
        ("お客様サービスセンターの名称", None, None, "フッター右上の見出し。貴社の窓口名"),
        ("電話番号", None, None, "現ひな型は 0120-000-000。貴社の受注・問い合わせ番号"),
        ("受付時間", None, None, "現ひな型は平日9:00-17:00（土日祝・弊社休業日を除く）。貴社の営業日・休業日"),
        ("会社概要リンク", None, None, "フッターに出すか"),
        ("ご利用ガイドリンク", None, None, "フッターに出すか"),
        ("ご利用規約リンク", None, None, "フッターに出すか"),
        ("よくあるご質問リンク", None, None, "フッターに出すか"),
        ("プライバシーポリシーリンク", None, None, "フッターに出すか"),
        ("セキュリティポリシーリンク", None, None, "方針が無い場合はリンク自体を外します"),
        ("特定商取引法に基づく表記リンク", None, None, "09シートの要否判断に合わせてください"),
        ("お問い合わせリンク", None, None, "フッターに出すか"),
    ],
    [22, 16, 42, 55],
    yesno_col=2,
)
page_setup(ws, "01 フッター")

# ---------------------------------------------------------------------------
# 02 会社概要
# ---------------------------------------------------------------------------
ws = wb.create_sheet("02_会社概要")
banner(
    ws, 4,
    "会社概要",
    "注文サイトに掲載する貴社の会社概要です。項目ごとに掲載する文言を記入してください。",
    "現ひな型はシステム提供会社（株式会社フライトソリューションズ）の設立・資本金・事業所が入っています。注文サイトの主体は貴社です。",
)
input_table(
    ws, 7,
    ["項目", "掲載要否", "ご記入内容", "現ひな型（掲載しません）"],
    [
        ("ページタイトル", "掲載する", "会社概要", "会社概要"),
        ("名称", "掲載する", None, "株式会社フライトソリューションズ"),
        ("設立", None, None, "1988年4月6日、ほか社名変更の沿革"),
        ("資本金", None, None, "1,205百万円"),
        ("代表者", None, None, "代表取締役 片山 圭一朗"),
        ("従業員数", None, None, "111名"),
        ("売上高", None, None, "144億円"),
        ("本社所在地", None, None, "東京都渋谷区恵比寿4-6-1"),
        ("本社電話・FAX", None, None, "03-3440-6100 / 03-5791-2241"),
        ("その他事業所", None, None, "仙台事業所が掲載されている"),
        ("会社ホームページURL", None, None, "https://www.flight.co.jp/"),
        ("追加したい項目", None, None, "事業内容、適格請求書発行事業者登録番号など。無ければ空欄"),
    ],
    [24, 16, 46, 42],
    yesno_col=2,
)
page_setup(ws, "02 会社概要")

# ---------------------------------------------------------------------------
# 03 お問い合わせ
# ---------------------------------------------------------------------------
ws = wb.create_sheet("03_お問い合わせ")
banner(
    ws, 4,
    "お問い合わせ",
    "問い合わせページに出す窓口と、フォームで受ける項目です。",
    "納期・出荷・返品・ログインは貴社CSが受ける業務です。宛先がシステム提供会社のままになると、問い合わせが貴社に届きません。",
)
input_table(
    ws, 7,
    ["項目", "掲載要否", "ご記入内容", "確認したいこと"],
    [
        ("表示する電話番号", None, None, "フッターと同じ番号でよいか"),
        ("表示するメールアドレス", None, None, "取引先に見せるアドレス"),
        ("フォームの送信先", None, None, "受信する貴社メールアドレス。表示アドレスと違ってよい"),
        ("受付時間の文言", None, None, "フッターと同じでよければ「フッターと同じ」"),
        ("用件の選択肢", None, None, "例：注文、納期、出荷、返品、ログイン、マスタ。必要な分類を列挙"),
        ("必須入力項目", None, None, "会社名、氏名、電話、メール、注文番号のうち必須のもの"),
        ("注文番号の入力", None, None, "納期・出荷・返品の問い合わせで注文番号を必須にするか"),
        ("回答目安", None, None, "例：2営業日以内。書けなければ空欄"),
        ("電話とメールの使い分け", None, None, "緊急時は電話、など案内文があれば記入"),
    ],
    [26, 16, 42, 48],
    yesno_col=2,
)
page_setup(ws, "03 お問い合わせ")

# ---------------------------------------------------------------------------
# 04 ご利用ガイド
# ---------------------------------------------------------------------------
ws = wb.create_sheet("04_ご利用ガイド")
banner(
    ws, 4,
    "ご利用ガイド",
    "標準の操作説明はそのまま使い、貴社固有のルールだけ文章を足していただきます。下表の「貴社で書いていただく内容」を記入してください。",
    "画面操作（検索、カート、ログイン）は標準機能で説明できます。注文モード、与信、承認、価格の見え方は貴社の業務ルールです。",
)
input_table(
    ws, 7,
    ["テーマ", "掲載要否", "貴社で書いていただく内容", "ガイドに反映するルール"],
    [
        ("ご利用開始まで", None, None, "初回のID通知、パスワード設定、誰が利用者を追加するか"),
        ("ログインと利用者", None, None, "1社何名までか、退職時の停止を誰が行うか"),
        ("商品の探し方", "掲載する", "標準の操作説明を使用", "カテゴリは貴社HPの分類を使用"),
        ("価格の見方", None, None, "上代・下代を誰に見せるか。代理店配下の店舗は下代を出さない"),
        ("在庫と入荷予定の見方", None, None, "入荷予定日はサイト一律で加算して表示する。加算日数が決まっていれば記入"),
        ("通常注文", None, None, "納品希望日が指定できる場合と、順次／揃い出ししか選べない場合"),
        ("カスタム注文", None, None, "デザインID、管理番号、仕様書の出し方。キャンセルはシステムで受け付けない"),
        ("展示会注文", None, None, "期間中のみボタンを出す。発注数の見え方。配下店舗には確保在庫を出さない"),
        ("代理注文", None, None, "代理店が配下店舗の分を注文できる範囲"),
        ("承認", None, None, "店舗注文を本社または代理店が承認するケースの有無と手順"),
        ("注文後の流れ", None, None, "与信、納期回答、出荷連絡がメールで届くこと"),
        ("購入履歴でできること", None, None, "再注文、納品書印刷、キャンセル依頼、返品依頼、仕様書の差し替え期限"),
        ("キャンセル", None, None, "出荷待ちまで、注文単位。基幹連携後は基幹側の可否に従う"),
        ("返品", None, None, "出荷後、明細単位。不可のときの扱い"),
        ("出荷先", None, None, "登録済み出荷先は取引先で追加・編集・削除できない。注文時の都度指定は可"),
        ("EDI注文", None, None, "画面では閲覧のみ。変更、キャンセル、返品は画面からできない"),
        ("操作説明（標準）", "掲載する", "標準の操作説明を使用", "会員登録、検索、カート、パスワード再設定など。修正がなければこの行のみで可"),
    ],
    [24, 16, 46, 52],
    yesno_col=2,
)
page_setup(ws, "04 ご利用ガイド")

# ---------------------------------------------------------------------------
# 05 ご利用規約
# ---------------------------------------------------------------------------
ws = wb.create_sheet("05_ご利用規約")
banner(
    ws, 4,
    "ご利用規約",
    "標準規約を土台に、貴社の取引で変える条項の結論を記入してください。全文を書き換える場合は別紙でも構いません。",
    "現ひな型は一般的な通信販売の規約です。本サイトは与信と承認があるBtoBなので、注文成立の時点を貴社のルールに合わせる必要があります。",
)
input_table(
    ws, 7,
    ["条項", "掲載要否", "貴社の結論", "決めていただきたいこと"],
    [
        ("規約の当事者", "掲載する", None, "利用者との契約当事者はローリングスジャパン合同会社でよいか"),
        ("適用範囲", None, None, "フロントサイトのみか。EDI注文にもこの規約を及ぼすか"),
        ("利用者登録", None, None, "登録できる人（取引先担当者、代理店、配下店舗）の範囲"),
        ("IDの管理", None, None, "IDの貸与禁止、退職者の削除責任をどちらが負うか"),
        ("注文の成立時点", None, None, "ボタン押下時か、与信可決時か、承認完了時か。カスタムは金額回答後か"),
        ("価格", None, None, "下代は取引先ごとに異なり、画面表示を正とするか、基幹を正とするか"),
        ("在庫・入荷予定", None, None, "表示在庫は確約か。入荷予定の日付加算がある旨を書くか"),
        ("承認が必要な注文", None, None, "承認前の注文の効力。否認時の扱い"),
        ("キャンセル", None, None, "出荷待ち・注文単位。基幹で不可の場合は成立したまま残る"),
        ("返品", None, None, "出荷後・明細単位。返品不可の条件があれば記入"),
        ("カスタム商品", None, None, "仕様確定後の変更、キャンセル不可を規約に書くか"),
        ("展示会", None, None, "期間外は注文できない。確保在庫の扱いを規約に書くか"),
        ("知的財産", "掲載する", None, "商品画像・仕様書の利用範囲。標準条項のままでよいか"),
        ("禁止事項", "掲載する", None, "標準条項に追加する禁止行為があれば記入"),
        ("免責", None, None, "連携遅延、在庫差異、メール不到達の責任分界"),
        ("準拠法・管轄", None, None, "合意管轄裁判所"),
        ("規約の変更", "掲載する", None, "サイト掲載による変更通知でよいか"),
    ],
    [24, 16, 44, 52],
    yesno_col=2,
)
page_setup(ws, "05 ご利用規約")

# ---------------------------------------------------------------------------
# 06 FAQ
# ---------------------------------------------------------------------------
ws = wb.create_sheet("06_よくあるご質問")
banner(
    ws, 5,
    "よくあるご質問",
    "取引先から繰り返し来る質問と、貴社として出す回答を記入してください。下の行は想定質問です。不要な行は掲載しないにして、足りなければ行を足してください。",
    "操作の一般論はガイドに任せ、このページは貴社固有の判断（誰が見えるか、いつキャンセルできるか）を短く答える場所です。",
)
faq_rows = [
    ("ログイン", "パスワードを忘れたときはどうすればよいか", None),
    ("利用者", "担当者を追加・削除したい", None),
    ("価格", "下代が表示されない", None),
    ("在庫", "入荷予定日はいつ届く日か", None),
    ("在庫", "展示会の確保数はどこで見るか", None),
    ("注文", "通常・カスタム・展示会の違い", None),
    ("注文", "納品日を指定できないことがある", None),
    ("注文", "順次発送と揃い出しの違い", None),
    ("代理注文", "代理店が店舗の代わりに注文できるか", None),
    ("承認", "注文したのに進まない", None),
    ("与信", "与信が通らなかったときの状態", None),
    ("納期", "納期回答はどこで確認するか", None),
    ("出荷", "出荷後の伝票番号はどこで見るか", None),
    ("キャンセル", "キャンセルできる期限", None),
    ("キャンセル", "キャンセルが拒否されたとき", None),
    ("返品", "返品できる単位と期限", None),
    ("カスタム", "仕様書はいつまで差し替えできるか", None),
    ("カスタム", "カスタム注文はキャンセルできるか", None),
    ("出荷先", "届ける場所を追加・変更できるか", None),
    ("EDI", "EDIで出した注文を画面で変更できるか", None),
    ("帳票", "納品書はどこから出すか", None),
    ("", "", None),
    ("", "", None),
    ("", "", None),
]
input_table(
    ws, 7,
    ["分類", "質問（想定。書き換えてよい）", "掲載要否", "回答", "補足"],
    [(a, b, None, c, None) for a, b, c in faq_rows],
    [16, 42, 14, 48, 28],
    yesno_col=3,
)
page_setup(ws, "06 よくあるご質問")

# ---------------------------------------------------------------------------
# 07 プライバシー
# ---------------------------------------------------------------------------
ws = wb.create_sheet("07_プライバシーポリシー")
banner(
    ws, 4,
    "プライバシーポリシー（個人情報保護方針）",
    "空欄になっている主体情報と、このサイトで実際に扱う個人情報の範囲を記入してください。条文の体裁は現ひな型を土台にできます。",
    "現ひな型は「株式会社○○○」「共同利用するグループ会社」が空欄です。取得項目に購買履歴が含まれるため、誰が管理するかを確定しないと掲載できません。",
)
input_table(
    ws, 7,
    ["項目", "掲載要否", "ご記入内容", "現ひな型で空欄または要確認の箇所"],
    [
        ("個人情報取扱事業者の正式名称", "掲載する", None, "文中の「当社」が指す会社。貴社名を想定"),
        ("個人情報保護管理者", None, None, "部署・役職・連絡先が空欄"),
        ("窓口の郵便・住所", None, None, "開示等の請求窓口"),
        ("窓口メールアドレス", None, None, "現ひな型はシステム提供側のアドレス。貴社窓口へ変更"),
        ("取得する情報", None, None, "氏名、所属、電話、メール、購買履歴、注文時に指定した出荷先"),
        ("利用目的", None, None, "受注、納期連絡、出荷連絡、請求、本サイトの提供。追加する目的があれば記入"),
        ("共同利用の有無", None, None, "グループ会社や物流委託先と共同利用するか。する場合は項目と目的"),
        ("委託先", None, None, "サイト運用、倉庫、配送。委託先の社名を掲載するか"),
        ("国外移転", None, None, "国外の委託や国外の親会社への提供があるか"),
        ("第三者提供", None, None, "基幹・WMS・EDIへ注文者情報を渡すことを第三者提供と整理するか"),
        ("開示・訂正・削除の方法", "掲載する", None, "請求先と、本人確認の方法"),
        ("手数料", None, None, "現ひな型は開示請求1件1,000円。踏襲するか"),
        ("安全管理", "掲載する", None, "標準の表現のままでよいか"),
        ("改定方法", "掲載する", None, "サイト掲載による周知でよいか"),
    ],
    [32, 14, 42, 48],
    yesno_col=2,
)
page_setup(ws, "07 プライバシーポリシー")

# ---------------------------------------------------------------------------
# 08 セキュリティ
# ---------------------------------------------------------------------------
ws = wb.create_sheet("08_セキュリティポリシー")
banner(
    ws, 4,
    "セキュリティポリシー",
    "見出しごとに掲載する本文を記入してください。方針をサイトに出さない場合は、掲載要否をすべて「掲載しない」にしてください。",
    "現ひな型は1〜11の見出しのみで、本文が空欄です。システム提供会社の方針を貴社サイトに掲載することはできません。",
)
secs = [
    "1. 目的",
    "2. セキュリティポリシーなどの公開",
    "3. 適用範囲",
    "4. 情報セキュリティ管理体制",
    "5. 情報セキュリティ対策の実施",
    "6. 情報セキュリティに関する社内規定の整備",
    "7. 情報セキュリティ教育の実施",
    "8. 適切な業務委託先管理の実施",
    "9. 法令などの遵守",
    "10. 情報セキュリティ内部監査の実施",
    "11. 継続的改善の実施",
]
input_table(
    ws, 7,
    ["見出し", "掲載要否", "本文", "補足"],
    [(s, None, None, "既存の方針文があれば貼り付け。無ければ空欄のまま「掲載しない」") for s in secs],
    [42, 16, 55, 36],
    yesno_col=2,
)
page_setup(ws, "08 セキュリティポリシー")

# ---------------------------------------------------------------------------
# 09 特商法
# ---------------------------------------------------------------------------
ws = wb.create_sheet("09_特定商取引法")
banner(
    ws, 4,
    "特定商取引法に基づく表記",
    "まず掲載するかどうかを決めてください。掲載する場合は、販売業者以下を貴社のBtoB条件で記入してください。",
    "現ひな型は消費者向け通販です。クレジット、コンビニ、代金引換、地域別送料、8日以内返品が書いてあり、本サイトの掛取引・与信・承認とは一致しません。",
)
input_table(
    ws, 7,
    ["項目", "掲載要否", "ご記入内容", "現ひな型（このまま使いません）"],
    [
        ("このページ自体を掲載するか", None, None, "取引先限定のBtoBでは掲載しない判断もあります。結論を記入"),
        ("販売業者", None, None, "株式会社フライトソリューションズ"),
        ("運営統括責任者", None, None, "片山 圭一朗"),
        ("所在地", None, None, "東京都渋谷区恵比寿4-6-1"),
        ("電話番号", None, None, "03-3440-6100"),
        ("メールアドレス", None, None, "システム提供側のアドレス"),
        ("サイトURL", None, None, "https://www.flight.co.jp/"),
        ("販売価格", None, None, "商品ごとに税別・税込のどちらで表示するか"),
        ("商品代金以外の必要料金", None, None, "送料、振込手数料、カスタムの追加費用"),
        ("送料", None, None, "地域別の消費者向け送料表が入っている。貴社の負担条件を記入"),
        ("支払方法", None, None, "クレジット、コンビニ、代金引換と書いてある。掛・請求書など貴社の方法"),
        ("支払時期", None, None, "注文後14日以内など消費者向け期限。締め・支払サイトを記入"),
        ("お届け方法", None, None, "配送事業者を書くか"),
        ("商品の引渡時期", None, None, "在庫品と入荷待ち、カスタム、展示会で時期が異なる"),
        ("注文の撤回・変更", None, None, "出荷待ちまでキャンセル可、カスタムは不可、など"),
        ("返品・交換", None, None, "到着後8日以内の消費者向け条件が入っている。貴社の返品条件"),
        ("不良品の取扱い", None, None, "送料負担をどちらが持つか"),
        ("資格・免許", None, None, "該当がなければ掲載しない"),
        ("サービス停止", None, None, "保守・障害時の停止を書くか"),
    ],
    [28, 16, 44, 48],
    yesno_col=2,
)
page_setup(ws, "09 特定商取引法")

wb.properties.title = "サイト掲載コンテンツのご提供依頼"
wb.properties.subject = "ローリングスジャパン様 BtoBサイト フッター掲載ページ"
wb.properties.creator = "フライトソリューションズ"

wb.save(OUT)
print("wrote", OUT)
